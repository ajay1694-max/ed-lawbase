"""export.py - DOCX exports for ED LawBase (python-docx if present, Markdown otherwise)."""
import io, re

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

NAVY = "1F3A5F"
CAUTION = "Research lead only. Cite from the certified copy / SCR / SCC and check later history before relying on it."


def _new(landscape=False):
    d = Document()
    sec = d.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.8))
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    return d


def _shade(cell, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(sh)


def _inline(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part.startswith("`"):
            par.add_run(part[1:-1]).font.name = "Consolas"
        else:
            par.add_run(part)


def _table(d, header, rows):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, NAVY)
    tr = t.rows[0]._tr.get_or_add_trPr()
    tr.append(OxmlElement("w:tblHeader"))
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _inline(cells[i].paragraphs[0], str(v or ""))
    return t


def _markdown(d, md):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip() or ln.strip() == "---":
            i += 1; continue
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
            body = [r for r in block[1:] if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
            _table(d, block[0], body)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            _inline(d.add_heading(level=len(m.group(1))), m.group(2))
        elif re.match(r"^\s*[-*]\s+", ln):
            _inline(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", ln))
        else:
            _inline(d.add_paragraph(), ln)
        i += 1


def _footer(d, attribution):
    p = d.add_paragraph()
    r = p.add_run(f"{CAUTION}\n{attribution}"); r.italic = True; r.font.size = Pt(8.5)


def _save(d):
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


# ------------------------------------------------------------------ public entry points
def citation_table(cases, attribution, title="Case law - citation table"):
    """cases: dicts with title, court, citation, decision_date, disposition, issues (labels), note"""
    if not HAVE_DOCX:
        md = [f"# {title}", "", "| Case | Court | Citation | Date | Disposition | Issues | Note |", "|---|---|---|---|---|---|---|"]
        md += [f"| {c['title']} | {c['court']} | {c.get('citation') or ''} | {c.get('decision_date') or ''} | "
               f"{c.get('disposition') or ''} | {'; '.join(c.get('issues', []))} | {c.get('note') or ''} |" for c in cases]
        return "\n".join(md + ["", CAUTION, attribution]).encode("utf-8"), "md"
    d = _new(landscape=True)
    d.add_heading(title, level=1)
    _table(d, ["#", "Case", "Court", "Citation", "Date", "Disposition", "Issues", "Note"],
           [(i + 1, c["title"], c["court"], c.get("citation"), c.get("decision_date"), c.get("disposition"),
             "; ".join(c.get("issues", [])), c.get("note")) for i, c in enumerate(cases)])
    _footer(d, attribution)
    return _save(d), "docx"


def judgment(meta, issues, curated, chunks, attribution):
    head = [meta["title"], f"{meta.get('court') or ''} | {meta.get('citation') or ''} | {meta.get('decision_date') or ''} | {meta.get('disposition') or ''}"]
    if not HAVE_DOCX:
        return ("\n\n".join([f"# {head[0]}", head[1], "Issues: " + "; ".join(issues)] + [c for c in chunks] + [CAUTION, attribution])).encode("utf-8"), "md"
    d = _new()
    d.add_heading(head[0], level=1)
    d.add_paragraph(head[1])
    if issues:
        _inline(d.add_paragraph(), "**Issues (auto/curated):** " + "; ".join(issues))
    for cur in curated:
        _inline(d.add_paragraph(), f"**Headnote ({cur.get('tier')}, {cur.get('stance') or 'stance not set'}):** {cur.get('summary') or ''}")
        if cur.get("body"):
            _markdown(d, cur["body"])
    d.add_heading("Text", level=2)
    for block in chunks:
        for para in re.split(r"\n\s*\n", block):
            para = re.sub(r"\[SECTION\]\s*#*\s*|\[TITLE\]\s*#*\s*", "", para).strip()
            if para and not re.fullmatch(r"[A-H]", para):
                d.add_paragraph(para)
    _footer(d, attribution)
    return _save(d), "docx"


def brief(meta, body, attribution):
    if not HAVE_DOCX:
        return (f"# {meta['title']}\n\n{body}\n\n{CAUTION}\n{attribution}").encode("utf-8"), "md"
    d = _new()
    d.add_heading(meta["title"], level=1)
    d.add_paragraph(f"{meta.get('date') or ''}  |  tier: {meta.get('tier')}  |  {meta.get('author') or ''}")
    _markdown(d, body)
    _footer(d, attribution)
    return _save(d), "docx"
